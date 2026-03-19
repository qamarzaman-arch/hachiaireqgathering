import os
import json
from datetime import datetime
from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt

def generate_pdf(session_data, output_path):
    """
    Generates a professional PDF requirements document.
    session_data: dict containing 'title', 'date', and 'steps' (list of dicts)
    """
    pdf = FPDF()
    pdf.add_page()

    # Title
    pdf.set_font("Arial", 'B', 24)
    pdf.set_text_color(67, 39, 118) # Hachiai Purple
    pdf.cell(0, 20, session_data.get('title', 'Requirements Document'), ln=True, align='C')

    # Metadata
    pdf.set_font("Arial", 'I', 10)
    pdf.set_text_color(128, 128, 128)
    pdf.cell(0, 10, f"Generated on: {session_data.get('date', datetime.now().strftime('%Y-%m-%d %H:%M'))}", ln=True, align='C')
    pdf.ln(10)

    # Overview
    pdf.set_font("Arial", 'B', 16)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 10, "1. Executive Overview", ln=True)
    pdf.set_font("Arial", '', 11)
    summary = session_data.get('summary', "This document outlines the workflow captured during the recording session.")
    pdf.multi_cell(0, 10, summary)
    pdf.ln(10)

    # Steps
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, "2. Workflow Steps", ln=True)
    pdf.ln(5)

    for i, step in enumerate(session_data.get('steps', []), 1):
        # Step Header
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, f"Step {i}: {step['title']}", ln=True)

        # Details
        pdf.set_font("Arial", '', 10)
        pdf.cell(30, 8, "Application:", border=0)
        pdf.set_font("Arial", 'B', 10)
        pdf.cell(0, 8, step.get('app_name', 'Unknown'), ln=True)

        pdf.set_font("Arial", 'I', 10)
        pdf.multi_cell(0, 8, step.get('description', 'No description provided.'))

        # Screenshot (if exists)
        img_path = step.get('screenshot_path')
        if img_path and os.path.exists(img_path):
            try:
                # Calculate width to fit page while maintaining aspect ratio
                pdf.image(img_path, x=20, w=170)
                pdf.ln(5)
            except Exception as e:
                pdf.set_text_color(255, 0, 0)
                pdf.cell(0, 10, f"[Error loading image: {str(e)}]", ln=True)
                pdf.set_text_color(0, 0, 0)

        pdf.ln(10)

        # Add page break if near bottom
        if pdf.get_y() > 250:
            pdf.add_page()

    pdf.output(output_path)
    return True

def generate_docx(session_data, output_path):
    """
    Generates a professional DOCX requirements document.
    """
    doc = Document()

    # Title
    title = doc.add_heading(session_data.get('title', 'Requirements Document'), 0)

    # Metadata
    p = doc.add_paragraph()
    p.add_run(f"Generated on: {session_data.get('date', datetime.now().strftime('%Y-%m-%d %H:%M'))}").italic = True

    # Overview
    doc.add_heading('1. Executive Overview', level=1)
    doc.add_paragraph(session_data.get('summary', "This document outlines the workflow captured during the recording session."))

    # Steps
    doc.add_heading('2. Workflow Steps', level=1)

    for i, step in enumerate(session_data.get('steps', []), 1):
        doc.add_heading(f"Step {i}: {step['title']}", level=2)

        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Property'
        hdr_cells[1].text = 'Value'

        row_cells = table.add_row().cells
        row_cells[0].text = 'Application'
        row_cells[1].text = step.get('app_name', 'Unknown')

        row_cells = table.add_row().cells
        row_cells[0].text = 'Description'
        row_cells[1].text = step.get('description', 'No description provided.')

        doc.add_paragraph() # Spacer

        img_path = step.get('screenshot_path')
        if img_path and os.path.exists(img_path):
            try:
                doc.add_picture(img_path, width=Inches(6))
            except Exception as e:
                doc.add_paragraph(f"[Error loading image: {str(e)}]").font.color.rgb = (255, 0, 0)

        doc.add_page_break()

    doc.save(output_path)
    return True

if __name__ == "__main__":
    import sys
    # Example usage for CLI testing
    if len(sys.argv) > 1:
        mock_data = {
            "title": "Purchase Order Process",
            "date": "2026-03-19 14:00",
            "summary": "This recording tracks the standard procedure for creating a purchase order in SAP and Excel.",
            "steps": [
                {"title": "Open SAP Logon", "description": "User launched the SAP GUI from the desktop.", "app_name": "SAP GUI", "screenshot_path": "test_img.webp"},
                {"title": "Enter Transaction ME21N", "description": "User navigated to the Create Purchase Order screen.", "app_name": "SAP GUI"}
            ]
        }

        action = sys.argv[1]
        out_path = sys.argv[2] if len(sys.argv) > 2 else "output.pdf"

        if action == "pdf":
            generate_pdf(mock_data, out_path)
        elif action == "docx":
            generate_docx(mock_data, out_path)
